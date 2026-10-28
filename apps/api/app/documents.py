import hashlib
import io
import re
import socket
import struct
import zipfile
from pathlib import PurePath
from typing import Annotated

from docx import Document as DocxDocument
from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile
from fastapi.responses import Response
from pydantic import BaseModel, Field
from pypdf import PdfReader
from sqlalchemy import and_, delete, exists, func, or_, select, text, true
from sqlalchemy.orm import Session

from app import storage
from app.audit import record
from app.auth import Principal, authorize_department, get_principal, require_permission
from app.core.config import get_settings
from app.db import get_db
from app.models import (
    Document,
    DocumentChunk,
    DocumentPermission,
    DocumentVersion,
    EmbeddingRecord,
    Job,
    Membership,
    uid,
    utcnow,
)

router = APIRouter(tags=["documents"])
MIMES = {".pdf": "application/pdf", ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document", ".txt": "text/plain", ".md": "text/markdown", ".csv": "text/csv"}


def document_filter(p: Principal, include_deleting: bool = False):
    allowed_depts = [d for d in p.department_ids if p.department_permissions is None or "documents:read" in p.department_permissions.get(d, set())]
    return and_(Document.deleted_at.is_(None), true() if include_deleting else Document.status != "deleting", Document.department_id.in_(allowed_depts), or_(Document.visibility == "department", Document.owner_id == p.id, exists(select(DocumentPermission.user_id).where(DocumentPermission.document_id == Document.id, DocumentPermission.user_id == p.id))))


def get_document(db: Session, p: Principal, document_id: str, include_deleting: bool = False,
                 *, for_update: bool = False) -> Document:
    require_permission(p, "documents:read")
    query = (select(Document).where(Document.id == document_id, document_filter(p, include_deleting))
             .execution_options(populate_existing=True))
    if for_update:
        query = query.with_for_update().execution_options(populate_existing=True)
    row = db.scalar(query)
    if not row:
        raise HTTPException(404, "Document not found")
    return row


def document_json(d):
    return {"id": d.id, "name": d.name, "department_id": d.department_id, "owner_id": d.owner_id, "mime_type": d.mime_type, "size": d.size, "visibility": d.visibility, "status": d.status, "scan_status": d.scan_status, "version": d.version, "embedding_status": d.embedding_status, "error_code": d.error_code, "created_at": d.created_at, "processed_at": d.processed_at}


def verify_file(name: str, mime: str | None, data: bytes) -> str:
    suffix = PurePath(name).suffix.lower()
    if suffix not in MIMES:
        raise HTTPException(415, "Supported file types: PDF, DOCX, TXT, Markdown, CSV")
    if not data or len(data) > get_settings().max_upload_mb * 1024 * 1024:
        raise HTTPException(413, "File is empty or exceeds the upload limit")
    if mime and mime.split(";")[0] not in {MIMES[suffix], "application/octet-stream", "text/plain" if suffix in {".md", ".csv"} else MIMES[suffix]}:
        raise HTTPException(415, "File MIME type does not match its extension")
    if suffix == ".pdf":
        if not data.startswith(b"%PDF-"):
            raise HTTPException(415, "Invalid PDF signature")
        if re.search(rb"/(JavaScript|JS|Launch|EmbeddedFile|RichMedia)\b", data):
            raise HTTPException(422, "Active or embedded PDF content is not accepted")
    elif suffix == ".docx":
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                entries = archive.infolist()
                names = set(archive.namelist())
                if "word/document.xml" not in names or "[Content_Types].xml" not in names or len(entries) > 2000:
                    raise ValueError("Invalid DOCX")
                if sum(e.file_size for e in entries) > 50 * 1024 * 1024:
                    raise ValueError("Oversized archive")
                for e in entries:
                    if e.filename.startswith("/") or ".." in PurePath(e.filename).parts or "vbaproject" in e.filename.lower() or "/embeddings/" in e.filename:
                        raise ValueError("Unsafe archive")
                    if e.filename.endswith(".rels") and b'TargetMode="External"' in archive.read(e):
                        raise ValueError("External links are not accepted")
        except (ValueError, zipfile.BadZipFile, RuntimeError) as exc:
            raise HTTPException(422, "Invalid or unsafe DOCX archive") from exc
    else:
        try:
            text = data.decode("utf-8-sig")
            if "\x00" in text or any(ord(c) < 9 for c in text):
                raise ValueError("Binary text")
        except (UnicodeError, ValueError) as exc:
            raise HTTPException(415, "Text files must be UTF-8 without binary content") from exc
    if b"EICAR-STANDARD-ANTIVIRUS-TEST-FILE" in data:
        raise HTTPException(422, "File quarantined by content scan")
    return MIMES[suffix]


def scan(data: bytes) -> str:
    s = get_settings()
    if not s.clamav_host:
        if s.require_virus_scan or s.is_production:
            raise ValueError("scanner_unavailable")
        return "structural_checks_only"
    with socket.create_connection((s.clamav_host, s.clamav_port), timeout=30) as sock:
        sock.sendall(b"zINSTREAM\x00")
        for i in range(0, len(data), 65536):
            chunk = data[i:i + 65536]
            sock.sendall(struct.pack("!I", len(chunk)) + chunk)
        sock.sendall(struct.pack("!I", 0))
        answer = sock.recv(4096)
        if b"OK" not in answer or b"FOUND" in answer:
            raise ValueError("malware_or_scan_error")
    return "clean"


def extract(data: bytes, mime: str) -> list[tuple[str, str]]:
    if mime == MIMES[".pdf"]:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted or len(reader.pages) > 500:
            raise ValueError("unsupported_pdf")
        pages = [(f"page {i + 1}", p.extract_text() or "") for i, p in enumerate(reader.pages)]
    elif mime == MIMES[".docx"]:
        doc = DocxDocument(io.BytesIO(data))
        pages = [(f"paragraph {i + 1}", p.text) for i, p in enumerate(doc.paragraphs) if p.text.strip()]
        pages += [(f"table {i + 1}", "\n".join(" | ".join(c.text for c in row.cells) for row in table.rows)) for i, table in enumerate(doc.tables)]
    else:
        pages = [("lines 1+", data.decode("utf-8-sig"))]
    if not any(text.strip() for _, text in pages):
        raise ValueError("no_extractable_text_requires_ocr")
    if sum(len(t) for _, t in pages) > 2_000_000:
        raise ValueError("extracted_text_limit")
    return pages


def chunks(pages: list[tuple[str, str]], size=1500, overlap=150):
    for location, content in pages:
        normalized = re.sub(r"[ \t]+", " ", content).strip()
        for offset in range(0, len(normalized), size - overlap):
            text = normalized[offset:offset + size]
            if text.strip():
                yield location + f" · chars {offset + 1}–{offset + len(text)}", text


def queue_ingestion(db, document, p):
    key = f"ingest:{document.id}:v{document.version}"
    job = db.scalar(select(Job).where(Job.idempotency_key == key).with_for_update()
                    .execution_options(populate_existing=True))
    if not job:
        job = Job(department_id=document.department_id, user_id=p.id, kind="ingest", resource_id=document.id, idempotency_key=key)
        db.add(job)
    else:
        if job.status in {"queued", "running", "retrying"}:
            raise HTTPException(409, "Document processing is already in progress")
        # Reindex runs with the requester's current permissions, even if the
        # original uploader has left. Document ownership and ACLs stay intact.
        job.user_id = p.id
        job.status, job.attempts, job.error_code = "queued", 0, None
        job.claim_token = None
        job.available_at = job.updated_at = utcnow()
    db.commit()
    from app.worker import dispatch_job
    dispatch_job(job.id)
    return job


@router.get("/documents")
def list_documents(p: Principal = Depends(get_principal), db: Session = Depends(get_db), department_id: str | None = None, q: str = "", limit: int = Query(30, ge=1, le=100), offset: int = Query(0, ge=0)):
    require_permission(p, "documents:read")
    query = select(Document).where(document_filter(p))
    if department_id:
        authorize_department(p, department_id, "documents:read")
        query = query.where(Document.department_id == department_id)
    if q:
        query = query.where(Document.name.ilike(f"%{q[:100]}%"))
    total = db.scalar(select(func.count()).select_from(query.subquery()))
    items = db.scalars(query.order_by(Document.created_at.desc()).limit(limit).offset(offset)).all()
    return {"items": [document_json(d) for d in items], "total": total, "limit": limit, "offset": offset}


@router.post("/documents/upload", status_code=202)
async def upload(file: Annotated[UploadFile, File()], department_id: Annotated[str, Form()], visibility: Annotated[str, Form()] = "department", document_id: Annotated[str | None, Form()] = None, p: Principal = Depends(get_principal), db: Session = Depends(get_db)):
    authorize_department(p, department_id, "documents:upload")
    if visibility not in {"department", "private"}:
        raise HTTPException(422, "Invalid visibility")
    data = await file.read(get_settings().max_upload_mb * 1024 * 1024 + 1)
    name = PurePath((file.filename or "document").replace("\\", "/")).name[:255]
    mime = verify_file(name, file.content_type, data)
    content_hash = hashlib.sha256(data).hexdigest()
    duplicate = db.scalar(select(Document).where(document_filter(p), Document.department_id == department_id, Document.content_hash == content_hash))
    if duplicate:
        raise HTTPException(409, "An accessible document with identical content already exists")
    if document_id:
        document = get_document(db, p, document_id, for_update=True)
        if document.owner_id != p.id or document.department_id != department_id:
            raise HTTPException(403, "Only the document owner may upload a new version")
        if document.status in {"queued", "processing"}:
            raise HTTPException(409, "Wait for the current document version to finish processing")
        document.version += 1
        document.status, document.error_code, document.extracted_text = "queued", None, ""
        # Fail closed: remove previous searchable chunks before the new version is processed.
        ids = select(DocumentChunk.id).where(DocumentChunk.document_id == document.id)
        db.execute(delete(EmbeddingRecord).where(EmbeddingRecord.chunk_id.in_(ids)))
        db.execute(delete(DocumentChunk).where(DocumentChunk.document_id == document.id))
    else:
        document = Document(id=uid(), department_id=department_id, owner_id=p.id, name=name, mime_type=mime, size=len(data), content_hash=content_hash, storage_key="", visibility=visibility, version=1)
        db.add(document)
    key = f"{department_id}/{document.id}/{document.version}/{uid()}"
    storage.put(key, data, mime)
    document.name, document.mime_type, document.size = name, mime, len(data)
    document.storage_key, document.content_hash = key, content_hash
    db.flush()
    db.add(DocumentVersion(document_id=document.id, department_id=department_id, version=document.version, content_hash=content_hash, storage_key=key))
    record(db, p, "document.upload", "document", document.id, department_id)
    queue_ingestion(db, document, p)
    db.refresh(document)
    return document_json(document)


@router.get("/documents/{document_id}")
def detail(document_id: str, p: Principal = Depends(get_principal), db: Session = Depends(get_db)):
    document = get_document(db, p, document_id)
    versions = db.scalars(select(DocumentVersion).where(DocumentVersion.document_id == document_id).order_by(DocumentVersion.version.desc())).all()
    result = {**document_json(document), "extracted_text": document.extracted_text[:100000], "versions": [{"version": v.version, "created_at": v.created_at, "content_hash": v.content_hash} for v in versions]}
    record(db, p, "document.read", "document", document_id, document.department_id)
    db.commit()
    return result


@router.get("/documents/{document_id}/status")
def status(document_id: str, p: Principal = Depends(get_principal), db: Session = Depends(get_db)):
    return document_json(get_document(db, p, document_id))


@router.get("/documents/{document_id}/download")
def download(document_id: str, p: Principal = Depends(get_principal), db: Session = Depends(get_db)):
    doc = get_document(db, p, document_id)
    if doc.status != "ready":
        raise HTTPException(409, "Document must complete scanning before download")
    record(db, p, "document.download", "document", doc.id, doc.department_id)
    db.commit()
    return Response(storage.get(doc.storage_key), media_type=doc.mime_type, headers={"Content-Disposition": "attachment; filename=download" + PurePath(doc.name).suffix, "X-Content-Type-Options": "nosniff", "Cache-Control": "no-store"})


@router.post("/documents/{document_id}/reindex", status_code=202)
def reindex(document_id: str, p: Principal = Depends(get_principal), db: Session = Depends(get_db)):
    doc = get_document(db, p, document_id, for_update=True)
    authorize_department(p, doc.department_id, "documents:upload")
    if doc.owner_id != p.id:
        authorize_department(p, doc.department_id, "documents:delete")
    if doc.status in {"queued", "processing"}:
        raise HTTPException(409, "Document processing is already in progress")
    doc.status, doc.error_code = "queued", None
    record(db, p, "document.reindex", "document", doc.id, doc.department_id)
    queue_ingestion(db, doc, p)
    return {"id": doc.id, "status": "queued"}


@router.delete("/documents/{document_id}")
def remove(document_id: str, p: Principal = Depends(get_principal), db: Session = Depends(get_db)):
    doc = get_document(db, p, document_id, include_deleting=True, for_update=True)
    if doc.owner_id != p.id:
        authorize_department(p, doc.department_id, "documents:delete")
    versions = list(db.scalars(select(DocumentVersion).where(DocumentVersion.document_id == doc.id)))
    # Make data inaccessible before object deletion; an object failure remains retriable.
    doc.status = "deleting"
    db.commit()
    for v in versions:
        storage.delete(v.storage_key)
    if db.bind is not None and db.bind.dialect.name == "postgresql":
        # The narrow database operation rechecks live authorization and removes
        # all source rows atomically while keeping the tombstone hidden by RLS.
        completed = db.scalar(text("SELECT public.finalize_document_deletion(:document_id)"),
                              {"document_id": doc.id})
        if not completed:
            raise HTTPException(403, "Document deletion is no longer authorized")
    else:
        ids = select(DocumentChunk.id).where(DocumentChunk.document_id == doc.id)
        db.execute(delete(EmbeddingRecord).where(EmbeddingRecord.chunk_id.in_(ids)))
        db.execute(delete(DocumentChunk).where(DocumentChunk.document_id == doc.id))
        db.execute(delete(DocumentVersion).where(DocumentVersion.document_id == doc.id))
        db.execute(delete(DocumentPermission).where(DocumentPermission.document_id == doc.id))
        doc.deleted_at, doc.extracted_text, doc.storage_key, doc.status = utcnow(), "", "", "deleted"
    record(db, p, "document.delete", "document", doc.id, doc.department_id)
    db.commit()
    return {"status": "deleted"}


class GrantInput(BaseModel):
    user_id: str = Field(max_length=36)


@router.put("/documents/{document_id}/permissions")
def grant(document_id: str, body: GrantInput, p: Principal = Depends(get_principal), db: Session = Depends(get_db)):
    doc = get_document(db, p, document_id)
    if doc.owner_id != p.id:
        raise HTTPException(403, "Only the owner can share a private document")
    if not db.get(Membership, (body.user_id, doc.department_id)):
        raise HTTPException(422, "The recipient must belong to the document department")
    if not db.get(DocumentPermission, (doc.id, body.user_id)):
        db.add(DocumentPermission(document_id=doc.id, user_id=body.user_id, owner_id=doc.owner_id))
    record(db, p, "document.share", "document", doc.id, doc.department_id)
    db.commit()
    return {"status": "granted"}


@router.delete("/documents/{document_id}/permissions/{user_id}")
def revoke(document_id: str, user_id: str, p: Principal = Depends(get_principal), db: Session = Depends(get_db)):
    doc = get_document(db, p, document_id)
    if doc.owner_id != p.id:
        raise HTTPException(403, "Only the owner can change document sharing")
    db.execute(delete(DocumentPermission).where(DocumentPermission.document_id == doc.id, DocumentPermission.user_id == user_id))
    record(db, p, "document.unshare", "document", doc.id, doc.department_id)
    db.commit()
    return {"status": "revoked"}
